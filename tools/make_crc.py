from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

rows = [
    ("DocumentActor", "Зберігає ідентифікатор та ім'я ролі; надає спільний final-метод відображення ролі.", "Customer, Translator, Editor"),
    ("Customer", "Завантажує документ, обирає мовну пару, отримує результат.", "TranslationProject, SourceDocument, LanguagePair"),
    ("Translator", "Перекладає сегменти та використовує терміни глосарія.", "TranslationProject, Glossary, TranslationSegment"),
    ("Editor", "Вичитує переклад, перевіряє якість і фіксує правки.", "TranslationProject, TranslationSegment, QualityControllable"),
    ("TranslationProject", "Керує документом, сегментами, станом і переходом до доставки.", "SourceDocument, Glossary, LanguagePair, TranslationSegment"),
    ("SourceDocument", "Описує файл оригіналу та його формат.", "TranslationProject, Customer"),
    ("LanguagePair", "Зберігає мову оригіналу та мову перекладу.", "TranslationProject, Customer"),
    ("Glossary", "Зберігає терміни та повертає рекомендований відповідник.", "TranslationProject, Translator"),
    ("TranslationSegment", "Зберігає номер, оригінал, переклад і стан перевірки сегмента.", "TranslationProject, Translator, Editor"),
]

wb = Workbook()
ws = wb.active
ws.title = "CRC-картки"
ws.append(["Клас", "обов'язки класу", "з ким співпрацює"])
for row in rows:
    ws.append(row)
for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
ws.column_dimensions["A"].width = 25
ws.column_dimensions["B"].width = 72
ws.column_dimensions["C"].width = 60
ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions
wb.save("docs/crc.xlsx")
